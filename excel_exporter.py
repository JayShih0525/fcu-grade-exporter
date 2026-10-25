from pathlib import Path
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import re


BASE_DIR = Path(__file__).resolve().parent

HTML_DIR = BASE_DIR / "semester_html"

OUTPUT_FILE = BASE_DIR / "fcu_grades.xlsx"

HEADERS = [
    "學年度",
    "學期",
    "課程編碼",
    "科目",
    "成績",
    "等第",
    "學分",
    "修習別",
    "通識類別"
]


def get_semester_info(filepath):
    """
    semester_111_1.html -> (111, 1)
    semester_111_2.html -> (111, 2)
    """

    match = re.fullmatch(
        r"semester_(\d+)_(\d+)\.html",
        filepath.name
    )

    if not match:
        return None

    return int(match.group(1)), int(match.group(2))


def extract_grade_table(filepath):
    """
    擷取 MyFCU 成績表格中的課程資料。
    """

    html = filepath.read_text(encoding="utf-8")

    soup = BeautifulSoup(html, "html.parser")

    # 尋找真正的成績表格
    table = soup.select_one(
        "table.table-bordered.table-condensed"
    )

    if table is None:
        print(f"警告：{filepath.name} 找不到成績表格")
        return []

    rows = []

    # 只擷取 AngularJS ng-repeat 產生的課程列
    course_rows = table.select(
        'tbody tr[ng-repeat="item in slistCtrl.semesterScoreList.list"]'
    )

    for tr in course_rows:

        cells = tr.find_all("td", recursive=False)

        if len(cells) != 9:
            continue

        values = [
            cell.get_text(" ", strip=True)
            for cell in cells
        ]

        # 轉換資料型態
        try:
            values[0] = int(values[0])
            values[1] = int(values[1])

            if values[4]:
                values[4] = float(values[4])

            if values[6]:
                values[6] = float(values[6])

        except ValueError:
            pass

        rows.append(values)

    return rows


def format_excel(ws):

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    widths = {
        "A": 12,
        "B": 10,
        "C": 18,
        "D": 42,
        "E": 12,
        "F": 12,
        "G": 12,
        "H": 14,
        "I": 20
    }

    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2):

        for cell in row:

            cell.alignment = Alignment(
                vertical="center"
            )

        # 成績顯示一位小數
        row[4].number_format = "0.0"

        # 學分最多顯示一位小數
        row[6].number_format = "0.#"


def export_excel():

    print("\n開始整理歷年學期成績...")

    if not HTML_DIR.exists():
        print(f"找不到資料夾：{HTML_DIR}")
        return

    html_files = []

    for filepath in HTML_DIR.glob("semester_*.html"):

        info = get_semester_info(filepath)

        if info:
            html_files.append((info, filepath))

    # 學年度、學期由小到大
    html_files.sort(key=lambda item: item[0])

    if not html_files:
        print("找不到 semester_111_1.html 格式的檔案")
        return

    wb = Workbook()

    ws = wb.active
    ws.title = "歷年學期成績"

    ws.append(HEADERS)

    total_courses = 0

    for (year, semester), filepath in html_files:

        print(f"\n處理：{filepath.name}")

        rows = extract_grade_table(filepath)

        print(f"找到 {len(rows)} 門課程")

        for row in rows:
            ws.append(row)

        total_courses += len(rows)

    format_excel(ws)

    wb.save(OUTPUT_FILE)

    print("\n" + "=" * 60)
    print("Excel 匯出完成")
    print("=" * 60)

    print(f"學期數量：{len(html_files)}")
    print(f"課程總數：{total_courses}")
    print(f"輸出位置：{OUTPUT_FILE}")


if __name__ == "__main__":
    export_excel()